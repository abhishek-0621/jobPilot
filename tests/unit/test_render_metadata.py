"""Rendered files must not name the generator (Chromium/Skia/python-docx) or a model."""

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from jobpilot.documents.render import (
    fit_one_page,
    html_to_pdf,
    render_resume_docx,
    render_resume_html,
    scrub_pdf_metadata,
)
from tests.unit.test_render import _doc, browser

MARKERS = ("chromium", "skia", "headlesschrome", "python-docx", "pypdf", "gemma", "ollama",
           "claude", "gpt")


def _pdf_meta_text(path: Path) -> str:
    r = PdfReader(path)
    info = {k: str(v) for k, v in (r.metadata or {}).items()}
    xmp = r.trailer["/Root"].get("/Metadata")
    return (repr(info) + (xmp.get_object().get_data().decode("latin-1") if xmp else "")).lower()


def _assert_clean(text: str) -> None:
    for m in MARKERS:
        assert m not in text, m


def test_scrub_pdf_metadata(tmp_path: Path) -> None:
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.add_metadata({"/Producer": "Skia/PDF m120", "/Creator": "Mozilla HeadlessChrome",
                    "/Title": "gemma4 output"})
    p = tmp_path / "x.pdf"
    w.write(p)
    scrub_pdf_metadata(p, author="Asha Example")
    _assert_clean(_pdf_meta_text(p))
    meta = PdfReader(p).metadata or {}
    assert meta.get("/Author") == "Asha Example"
    assert len(PdfReader(p).pages) == 1


def test_docx_metadata_is_candidate(tmp_path: Path) -> None:
    out = tmp_path / "r.docx"
    render_resume_docx(_doc(), out)
    cp = Document(str(out)).core_properties
    assert cp.author == cp.last_modified_by == "Asha Example"
    assert (cp.title, cp.comments, cp.keywords, cp.subject, cp.category) == ("",) * 5
    blob = repr([getattr(cp, a) for a in dir(cp) if not a.startswith("_")]).lower()
    _assert_clean(blob)
    import zipfile
    with zipfile.ZipFile(out) as z:
        for name in ("docProps/core.xml", "docProps/app.xml"):
            if name in z.namelist():
                _assert_clean(z.read(name).decode("utf-8").lower())


def test_resume_layout_has_no_typographic_dashes(tmp_path: Path) -> None:
    html = render_resume_html(_doc())
    assert "—" not in html and "–" not in html
    out = tmp_path / "r.docx"
    render_resume_docx(_doc(), out)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "—" not in text and "–" not in text


@browser
@pytest.mark.browser
async def test_real_pdfs_have_clean_metadata(tmp_path: Path) -> None:
    doc, _ = await fit_one_page(_doc(), tmp_path / "r.pdf")
    _assert_clean(_pdf_meta_text(tmp_path / "r.pdf"))
    assert (PdfReader(tmp_path / "r.pdf").metadata or {}).get("/Author") == doc.name
    await html_to_pdf("<p>Hello</p>", tmp_path / "c.pdf", author="Asha Example")
    _assert_clean(_pdf_meta_text(tmp_path / "c.pdf"))
