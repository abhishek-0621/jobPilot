from pathlib import Path

import pytest
from docx import Document

from jobpilot.contracts.enums import ErrorCode
from jobpilot.data.tables import CandidateFactRow
from jobpilot.documents.model import Bullet, ResumeDoc, RoleBlock, base_resume
from jobpilot.documents.render import (
    DocRenderError,
    chromium_available,
    fit_one_page,
    html_to_pdf,
    render_resume_docx,
    render_resume_html,
)
from jobpilot.profile.importer import _flatten
from jobpilot.profile.schema import load_profile_file

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "profile_example.yaml"
browser = pytest.mark.skipif(not chromium_available(), reason="Chromium not installed")


def _facts(unverified: set[str] = frozenset()) -> list[CandidateFactRow]:
    pf = load_profile_file(FIXTURE)
    return [CandidateFactRow(fact_id=f.fact_id, profile_version=1, category=f.category,
                             field=f.field, value=f.value,
                             verified_by_user=f.fact_id not in unverified)
            for f in _flatten(pf)]


def _doc(unverified: set[str] = frozenset()) -> ResumeDoc:
    return base_resume(load_profile_file(FIXTURE), _facts(unverified))


def test_base_resume_uses_only_verified_facts() -> None:
    doc = _doc({"fact_exp_002", "fact_skill_genai", "fact_ach_001"})
    assert doc.name == "Asha Example"
    assert doc.headline == "AI Engineer"
    assert doc.contact_line.startswith("+919000000000 | asha@example.com | ")
    role = doc.roles[0]
    assert (role.employer, role.title, role.start, role.end) == (
        "Example Corp", "AI Engineer", "2023-01", "present")
    assert [b.fact_id for b in role.bullets] == ["fact_exp_001"]
    assert [p.title for p in doc.projects] == ["Example RAG Assistant"]
    assert doc.skills == [("Languages", ["Python", "SQL"])]
    assert doc.achievements == []
    assert doc.education == ["M.Tech Computer Science, Example Institute of Technology (2022)"]
    assert doc.summary == ""


def test_html_escapes_fact_text() -> None:
    doc = _doc()
    doc.roles[0].bullets[0].text = "<script>alert(1)</script>"
    html = render_resume_html(doc)
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


@browser
@pytest.mark.browser
async def test_pdf_single_page_for_fixture(tmp_path: Path) -> None:
    assert await html_to_pdf(render_resume_html(_doc()), tmp_path / "r.pdf") == 1


def _big_doc(n_roles: int, per_role: int, words: int = 12) -> ResumeDoc:
    roles, prio = [], 0
    for r in range(n_roles):
        bullets = []
        for b in range(per_role):
            bullets.append(Bullet(fact_id=f"f{r}_{b}", priority=prio,
                                  text=" ".join(["Synthetic"] * words) + f" {r}.{b}"))
            prio += 1
        roles.append(RoleBlock(employer=f"Example Co {r}", title="Engineer", start="2020",
                               end="2021", location=None, bullets=bullets))
    return ResumeDoc(name="Asha Example", headline="Engineer", contact_line="asha@example.com",
                     summary="", skills=[], roles=roles, projects=[], education=[],
                     achievements=[])


@browser
@pytest.mark.browser
async def test_overflow_drops_low_priority_bullets(tmp_path: Path) -> None:
    doc = _big_doc(n_roles=4, per_role=10)
    final, dropped = await fit_one_page(doc, tmp_path / "r.pdf", max_drops=40)
    assert dropped
    all_prios = {b.fact_id: b.priority for r in doc.roles for b in r.bullets}
    kept = {b.fact_id for r in final.roles for b in r.bullets}
    assert min(all_prios[d] for d in dropped) > max(
        all_prios[k] for k in kept if not k.endswith("_0"))
    assert all(f"f{r}_0" in kept for r in range(4))
    assert await html_to_pdf(render_resume_html(final), tmp_path / "x.pdf") == 1


def test_docx_contains_all_sections(tmp_path: Path) -> None:
    doc = _doc().model_copy(update={"summary": "Synthetic summary."})
    out = tmp_path / "r.docx"
    render_resume_docx(doc, out)
    d = Document(str(out))
    assert not d.tables
    heads = [p.text for p in d.paragraphs if p.style.name == "Heading 2"]
    assert heads == ["Summary", "Skills", "Experience", "Projects", "Education", "Achievements"]
    assert d.paragraphs[0].text == "Asha Example"
    bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
    assert "Won an internal hackathon" in bullets


@browser
@pytest.mark.browser
async def test_unfittable_raises_doc_render(tmp_path: Path) -> None:
    doc = _big_doc(n_roles=30, per_role=2, words=30)
    with pytest.raises(DocRenderError) as e:
        await fit_one_page(doc, tmp_path / "r.pdf")
    assert e.value.code == ErrorCode.DOC_RENDER
