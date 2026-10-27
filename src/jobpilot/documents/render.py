"""Render a ResumeDoc to HTML, PDF (Playwright Chromium) and DOCX; fit the PDF to one page."""

from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from functools import cache
from pathlib import Path
from typing import Any

from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape
from playwright.async_api import async_playwright
from playwright.sync_api import sync_playwright
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject

from jobpilot.contracts.enums import ErrorCode
from jobpilot.documents.model import ResumeDoc, RoleBlock

_TEMPLATES = Path(__file__).resolve().parent / "templates"
_env = Environment(loader=FileSystemLoader(_TEMPLATES),
                   autoescape=select_autoescape(["html", "j2"]))
_MARGIN = {side: "12mm" for side in ("top", "right", "bottom", "left")}


class DocRenderError(Exception):
    code = ErrorCode.DOC_RENDER

    def __init__(self, detail: str = "") -> None:
        super().__init__(f"{self.code.value}: {detail}" if detail else self.code.value)


@cache
def chromium_available() -> bool:
    """Import-time only (cached; uses the sync API): for pytest skip markers, never
    call it from inside a running event loop."""
    try:
        with sync_playwright() as p:
            return Path(p.chromium.executable_path).exists()
    except Exception:
        return False


def render_resume_html(doc: ResumeDoc) -> str:
    return _env.get_template("resume.html.j2").render(doc=doc)


def render_cover_letter_html(name: str, greeting: str, paragraphs: list[str],
                             closing: str) -> str:
    return _env.get_template("cover_letter.html.j2").render(
        name=name, greeting=greeting, paragraphs=paragraphs, closing=closing)


@asynccontextmanager
async def _browser() -> AsyncIterator[Any]:
    """One headless Chromium; launch failures surface as DocRenderError."""
    try:
        async with async_playwright() as p:
            try:
                browser = await p.chromium.launch(headless=True)
            except Exception as e:
                raise DocRenderError(f"chromium launch failed: {type(e).__name__}") from e
            try:
                yield browser
            finally:
                await browser.close()
    except DocRenderError:
        raise
    except Exception as e:
        raise DocRenderError(f"playwright failed: {type(e).__name__}") from e


def scrub_pdf_metadata(path: Path, author: str | None = None) -> None:
    """Replace the document info (Chromium writes Creator "HeadlessChrome" and Producer
    "Skia/PDF") with the candidate's name only, and drop any XMP stream."""
    writer = PdfWriter(clone_from=PdfReader(path))
    name = author or ""
    writer.metadata = None
    writer.add_metadata({"/Author": name, "/Creator": name, "/Producer": name})
    root = writer._root_object
    if NameObject("/Metadata") in root:
        del root[NameObject("/Metadata")]
    tmp = path.with_name(path.name + ".tmp")
    try:
        with tmp.open("wb") as fh:
            writer.write(fh)
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


async def _print_pdf(browser: Any, html: str, out: Path, author: str | None = None) -> int:
    out.parent.mkdir(parents=True, exist_ok=True)
    page = await browser.new_page()
    try:
        await page.set_content(html, wait_until="load")
        await page.pdf(path=str(out), format="A4", margin=_MARGIN, print_background=False)
    finally:
        await page.close()
    scrub_pdf_metadata(out, author)
    return len(PdfReader(out).pages)


async def html_to_pdf(html: str, out: Path, author: str | None = None) -> int:
    async with _browser() as browser:
        return await _print_pdf(browser, html, out, author)


def _drop_one(doc: ResumeDoc) -> str | None:
    """Remove the highest-priority-number droppable bullet; ties go to the later block."""
    blocks = [r.bullets for r in doc.roles] + [r.bullets for r in doc.projects]
    candidates = [(b.priority, bi, i, b.fact_id)
                  for bi, bl in enumerate(blocks) for i, b in enumerate(bl) if i > 0]
    candidates += [(b.priority, len(blocks), i, b.fact_id)
                   for i, b in enumerate(doc.achievements)]
    if not candidates:
        return None
    _, bi, i, fact_id = max(candidates)
    (blocks[bi] if bi < len(blocks) else doc.achievements).pop(i)
    return fact_id


async def fit_one_page(doc: ResumeDoc, out: Path,
                       max_drops: int = 6) -> tuple[ResumeDoc, list[str]]:
    """Render to ``out``, dropping low-priority bullets until one page; reuses one
    browser across iterations and never leaves an over-long ``out`` behind."""
    doc = doc.model_copy(deep=True)
    dropped: list[str] = []
    async with _browser() as browser:
        while await _print_pdf(browser, render_resume_html(doc), out, doc.name) != 1:
            fact_id = _drop_one(doc) if len(dropped) < max_drops else None
            if fact_id is None:
                out.unlink(missing_ok=True)
                reason = (f"still over one page after dropping {len(dropped)} bullets"
                          if len(dropped) >= max_drops
                          else "over one page and no droppable bullets left")
                raise DocRenderError(reason)
            dropped.append(fact_id)
    return doc, dropped


def _blocks_docx(d: Document, items: list[RoleBlock]) -> None:
    for r in items:
        head = r.title + (f", {r.employer}" if r.employer else "")
        head += f" | {r.location}" if r.location else ""
        dates = r.start + (f" - {r.end}" if r.end else "")
        p = d.add_paragraph()
        p.add_run(head).bold = True
        if dates:
            p.add_run(f"\t{dates}")
        for b in r.bullets:
            d.add_paragraph(b.text, style="List Bullet")


def render_resume_docx(doc: ResumeDoc, out: Path) -> None:
    d = Document()
    d.add_heading(doc.name, level=1)
    d.add_paragraph(doc.headline)
    d.add_paragraph(doc.contact_line)
    if doc.summary:
        d.add_heading("Summary", level=2)
        d.add_paragraph(doc.summary)
    if doc.skills:
        d.add_heading("Skills", level=2)
        for group, items in doc.skills:
            p = d.add_paragraph()
            p.add_run(f"{group}: ").bold = True
            p.add_run(", ".join(items))
    if doc.roles:
        d.add_heading("Experience", level=2)
        _blocks_docx(d, doc.roles)
    if doc.projects:
        d.add_heading("Projects", level=2)
        _blocks_docx(d, doc.projects)
    if doc.education:
        d.add_heading("Education", level=2)
        for e in doc.education:
            d.add_paragraph(e)
    if doc.achievements:
        d.add_heading("Achievements", level=2)
        for b in doc.achievements:
            d.add_paragraph(b.text, style="List Bullet")
    cp = d.core_properties  # the default template names python-docx; use the candidate
    cp.author = cp.last_modified_by = doc.name
    for attr in ("title", "subject", "keywords", "comments", "category",
                 "content_status", "identifier", "language", "version"):
        setattr(cp, attr, "")
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
